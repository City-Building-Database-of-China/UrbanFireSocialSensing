"""Scan public tables, structured text and documentation for sensitive content.

Location-field names in a data dictionary or explanatory prose are definitions,
not location values. Column headers, JSON keys and actual content are checked
separately, preserving explicitly anonymous IDs and boolean availability flags.
"""
import argparse
import csv
import gzip
import json
from pathlib import Path
import re

ROOT = Path(__file__).resolve().parents[2]
SKIP = {".git", ".venv", "venv", "__pycache__"}
FORMATS = {".csv", ".xlsx", ".json", ".parquet", ".txt", ".md", ".cff", ".py"}
BLOCKED = {"longitude", "latitude", "lng", "lat", "lon", "xcoordinate", "ycoordinate",
    "address", "exactaddress", "aoi", "aoiname", "aoitype", "community", "communityname",
    "exactpoi", "poi", "poiname", "poiid", "uid", "userid", "username", "weiboid", "mid",
    "url", "weibourl", "sourceurl", "buildingid", "originalbuildingid", "recordid", "caseid",
    "geometry", "geom", "wkt", "coordinates", "rawtext", "rawweibotext", "weibotext",
    "originaltext", "posttext", "mtxt", "cleanmtxt", "crosswalk", "bdoid", "buildingname"}
ALLOWED_IDS = {"publicbuildingid", "publiceventid", "publicrecordid", "sampleid",
               "hascoordinates", "hasplacename", "hasaddress"}
FULL_FEATURES = {"Projected_Area", "Height", "Plot_Ratio", "Effective_Age", "Mean_Age_Score",
                 "Pop_Density", "Price", "Fee", "EUI"}
LOCAL_PATH = re.compile(r"(?<![A-Za-z0-9])(?:[A-Za-z]:[\\/]|/(?:Users|home|mnt|tmp)/|\\\\[A-Za-z0-9_.-]+\\)")
WEIBO_LINK = re.compile(r"(?:https?://)?(?:[A-Za-z0-9-]+\.)?(?:weibo\.com|weibo\.cn|t\.cn)/\S+", re.I)
COORDINATE_VALUE = re.compile(r"(?:longitude|latitude|\blng\b|\blat\b)\s*[=:]\s*[-+]?\d+\.\d{3,}", re.I)
SPATIAL_SUFFIXES = {".shp", ".shx", ".dbf", ".gpkg", ".geojson", ".kml", ".kmz"}


def normalized(field):
    return re.sub(r"[^a-z0-9]", "", str(field).lower())


def files_in(root):
    return (p for p in sorted(root.rglob("*")) if p.is_file() and not any(x in SKIP for x in p.relative_to(root).parts))


def scan(root=ROOT):
    root = Path(root).resolve()
    findings, scanned = [], 0

    def fail(p, reason):
        findings.append({"file": p.relative_to(root).as_posix(), "reason": reason})

    def check_fields(p, fields):
        for field in fields:
            key = normalized(field)
            if key in BLOCKED and key not in ALLOWED_IDS:
                fail(p, f"Sensitive data field: {field}")

    def check_content(p, value, prose=False, source_code=False):
        value = str(value)
        if LOCAL_PATH.search(value):
            fail(p, "Absolute workstation path")
        if WEIBO_LINK.search(value):
            fail(p, "Social-media record URL")
        if not source_code and COORDINATE_VALUE.search(value):
            fail(p, "Explicit precise coordinate value")
        # Long Chinese narrative content is not part of the released aggregate schemas.
        if not prose and len(value) > 180 and len(re.findall(r"[\u4e00-\u9fff]", value)) > 60:
            fail(p, "Long narrative content requires removal from the public data")

    def check_json(p, value):
        if isinstance(value, dict):
            check_fields(p, value.keys())
            for v in value.values():
                check_json(p, v)
        elif isinstance(value, list):
            for v in value:
                check_json(p, v)
        else:
            check_content(p, value)

    def check_rows(p, headers, rows):
        headers = list(headers)
        check_fields(p, headers)
        count = 0
        for row in rows:
            count += 1
            for value in row:
                if isinstance(value, str):
                    check_content(p, value)
        if FULL_FEATURES.issubset(headers) and count > 10000:
            fail(p, "A row-level matrix larger than the authorized Shanghai demo")

    for path in files_in(root):
        if path.suffix.lower() in SPATIAL_SUFFIXES:
            fail(path, "Precise spatial file type is outside the public scope")
            continue
        if 'crosswalk' in path.name.lower():
            fail(path, "Identity/location crosswalk")
        if path.suffix.lower() not in FORMATS and not path.name.endswith('.csv.gz'):
            continue
        scanned += 1
        try:
            if path.suffix.lower() == '.csv' or path.name.endswith('.csv.gz'):
                opener = gzip.open if path.name.endswith('.gz') else open
                with opener(path, 'rt', encoding='utf-8-sig', newline='') as handle:
                    reader = csv.reader(handle)
                    check_rows(path, next(reader, []), reader)
            elif path.suffix.lower() == '.xlsx':
                from openpyxl import load_workbook
                book = load_workbook(path, read_only=True, data_only=False)
                for sheet in book:
                    rows = sheet.iter_rows(values_only=True)
                    check_rows(path, next(rows, ()), rows)
                book.close()
            elif path.suffix.lower() == '.parquet':
                import pandas as pd
                frame = pd.read_parquet(path)
                check_rows(path, frame.columns, frame.itertuples(index=False, name=None))
            elif path.suffix.lower() == '.json':
                check_json(path, json.loads(path.read_text(encoding='utf-8-sig')))
            else:
                check_content(path, path.read_text(encoding='utf-8-sig'), prose=True,
                              source_code=path.suffix.lower()=='.py')
        except Exception as exc:
            fail(path, f"Could not fully scan file ({type(exc).__name__}); review required")
    unique = list({(f['file'], f['reason']): f for f in findings}.values())
    return {"status": "passed" if not unique else "failed", "scanned_files": scanned,
            "findings": unique, "scan_scope": "Headers, JSON keys, all table rows, URLs, paths and explicit coordinates"}


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--root', type=Path, default=ROOT)
    p.add_argument('--report', type=Path)
    args = p.parse_args()
    result = scan(args.root)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(result, indent=2)+'\n', encoding='utf-8')
    print(json.dumps(result, indent=2))
    raise SystemExit(0 if result['status']=='passed' else 1)


if __name__ == '__main__':
    main()
